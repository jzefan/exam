from io import BytesIO
import uuid

import pytest
from docx import Document
from sqlalchemy import select

from app.auth.schemas import UserCreate
from app.auth.security import create_access_token
from app.auth.service import create_user
from app.papers.models import PaperImportSession
from app.papers.schemas import PaperImportRecognizeRequest
from app.rbac.models import Organization, Role
from app.questions.schemas import (
    QuestionImportDocumentRecognizeResponse,
    QuestionImportDocumentSummary,
    QuestionImportDraft,
)


def test_paper_import_recognize_request_accepts_pdf_page_data_urls():
    page_image = "data:image/jpeg;base64," + ("a" * 300_000)

    payload = PaperImportRecognizeRequest.model_validate(
        {
            "file_name": "paper.pdf",
            "source_format": "pdf",
            "raw_text": "[IMAGE:page-1]",
            "images": [
                {
                    "image_id": "page-1",
                    "url": page_image,
                    "order": 1,
                    "page": 1,
                    "alt": "第 1 页",
                }
            ],
        }
    )

    assert payload.images[0].url == page_image


def test_standard_paper_rule_split_keeps_all_section_questions_and_internal_numbering():
    from app.questions.service import _build_standard_paper_drafts_from_text

    parts: list[str] = ["一、单项选择题（每小题 2 分，共 50 分）"]
    for number in range(1, 26):
        stem = f"{number}.单选题{number}（ ）"
        if number == 25:
            stem += "\n1.内部概念模式说明\n2.内部物理模式说明\n3.内部外模式说明"
        parts.append(f"{stem}\nA. 甲 B. 乙 C. 丙 D. 丁\n正确答案： A")

    parts.append("二、多选题（每题 2 分，共 10 分）")
    for number in range(1, 6):
        parts.append(f"{number}.多选题{number}（ ）\nA. 甲 B. 乙 C. 丙 D. 丁\n正确答案： AB")

    parts.append("三、填空题（每题 2 分，共 10 分）")
    for number in range(1, 6):
        parts.append(f"{number}.填空题{number}是 。\n正确答案： 答案{number}")

    parts.append("四、判断题（每小题 1 分，共 5 分）")
    for number in range(1, 6):
        parts.append(f"{number}.判断题{number}。（ ）\n正确答案： 正确")

    parts.append("五、简答题（每小题 5 分，共 25 分）")
    for number in range(1, 6):
        if number == 2:
            parts.append(f"\n[IMG:stem-er.png]\n{number}.简答题{number}\n正确答案： 要点{number}")
        elif number == 3:
            parts.append(f"\n[IMG:answer-er.png]\n{number}.简答题{number}")
        else:
            parts.append(f"{number}.简答题{number}\n正确答案： 要点{number}")

    drafts = _build_standard_paper_drafts_from_text(
        "\n".join(parts),
        {
            "stem-er.png": "/api/uploads/files/stem-er.png",
            "answer-er.png": "/api/uploads/files/answer-er.png",
        },
    )

    assert len(drafts) == 45
    assert [draft.type.value for draft in drafts].count("choice") == 30
    assert [draft.type.value for draft in drafts].count("fill_in") == 5
    assert [draft.type.value for draft in drafts].count("true_false") == 5
    assert [draft.type.value for draft in drafts].count("short_answer") == 5
    assert drafts[0].options == {"A": "甲", "B": "乙", "C": "丙", "D": "丁"}
    assert "选择题选项不完整" not in drafts[0].issues
    assert drafts[25].options == {"A": "甲", "B": "乙", "C": "丙", "D": "丁"}
    assert "选择题选项不完整" not in drafts[25].issues
    assert "内部概念模式说明" in drafts[24].content_text
    assert drafts[41].images[0].url == "/api/uploads/files/stem-er.png"
    assert drafts[42].images == []
    assert drafts[42].answer_images[0].url == "/api/uploads/files/answer-er.png"
    assert "未识别到答案" not in drafts[42].issues


def test_choice_options_split_when_multiple_options_share_one_line():
    from app.questions.service import build_import_draft_from_segment

    drafts = [
        build_import_draft_from_segment(
            "7.DBMS指的是以下哪个选项（ ）\n"
            "A. 数据库系统 B. 数据库信息系统 C. 数据库管理系统 D. 数据库开发系统\n"
            "正确答案： C",
            type_hint="单项选择题",
        ),
        build_import_draft_from_segment(
            "8.下面选项中，能够实现查询表中记录的关键字是（ ）\n"
            "A. DROP B.SELECT C.UPDATE D.DELETE\n"
            "正确答案： B",
            type_hint="单项选择题",
        ),
        build_import_draft_from_segment(
            "9.SQL 中用于删除表的命令是（ ） A. SELECT B. DROP C. UPDATE D. INSERT\n"
            "正确答案： B",
            type_hint="单项选择题",
        ),
    ]

    assert drafts[0].content_text == "DBMS指的是以下哪个选项（ ）"
    assert drafts[0].options == {
        "A": "数据库系统",
        "B": "数据库信息系统",
        "C": "数据库管理系统",
        "D": "数据库开发系统",
    }
    assert drafts[1].content_text == "下面选项中，能够实现查询表中记录的关键字是（ ）"
    assert drafts[1].options == {"A": "DROP", "B": "SELECT", "C": "UPDATE", "D": "DELETE"}
    assert drafts[2].content_text == "SQL 中用于删除表的命令是（ ）"
    assert drafts[2].options == {"A": "SELECT", "B": "DROP", "C": "UPDATE", "D": "INSERT"}
    assert all("选择题选项不完整" not in draft.issues for draft in drafts)


def test_import_draft_exposes_multi_choice_for_glued_letter_answers():
    from app.papers.service import question_create_from_import_draft

    def make_draft(draft_id: str, answer_text: str) -> QuestionImportDraft:
        return QuestionImportDraft.model_validate(
            {
                "draft_id": draft_id,
                "raw_text": f"{draft_id}. 示例题",
                "title": "示例题",
                "type": "choice",
                "content_text": "示例题",
                "options": {"A": "甲", "B": "乙", "C": "丙", "D": "丁"},
                "answer_text": answer_text,
                "segment_source": "rule",
                "type_confidence": "high",
                "boundary_confidence": "high",
            }
        )

    single = question_create_from_import_draft(make_draft("single", "A"))
    multi_glued = question_create_from_import_draft(make_draft("multi-glued", "BD"))
    multi_separated = question_create_from_import_draft(make_draft("multi-separated", "A、C"))

    assert single.answer["correct"] == "A"
    assert multi_glued.answer["correct"] == ["B", "D"]
    assert multi_separated.answer["correct"] == ["A", "C"]


async def _teacher(db_session):
    org = Organization(name="Import Org", type="school", is_active=True)
    role = await db_session.scalar(select(Role).where(Role.name == "teacher"))
    if role is None:
        role = Role(name="teacher", display_name="Teacher", is_system=True)
        db_session.add(role)
    db_session.add(org)
    await db_session.flush()
    return await create_user(
        db_session,
        UserCreate(
            username="paper-import-teacher",
            email="paper-import-teacher@example.com",
            password="teacherpass123",
            full_name="Import Teacher",
            role_name="teacher",
            org_id=org.id,
        ),
    )


@pytest.mark.asyncio
async def test_paper_import_recognize_and_confirm_creates_paper(client, db_session, monkeypatch):
    teacher = await _teacher(db_session)
    client.headers.update({"Authorization": f"Bearer {create_access_token(teacher.id, '')}"})

    async def fake_request(_prompt: str) -> list[dict]:
        return [
            {
                "type": "choice",
                "content_text": "MySQL 属于哪类数据库？",
                "options": {"A": "关系型", "B": "缓存"},
                "answer_text": "A",
                "analysis": "",
                "difficulty": 3,
                "raw_text": "MySQL 属于哪类数据库？",
                "images": [],
            }
        ]

    monkeypatch.setattr("app.questions.service._request_doc_recognition_questions", fake_request)

    recognize_response = await client.post(
        "/api/papers/import/recognize",
        json={
            "file_name": "history.md",
            "source_format": "md",
            "root_knowledge_point_id": None,
            "raw_text": "1. 单选题 MySQL 属于哪类数据库？\nA. 关系型\nB. 缓存\n答案：A",
            "images": [],
        },
    )
    assert recognize_response.status_code == 200
    recognized = recognize_response.json()
    assert recognized["session_id"]
    assert recognized["drafts"]

    draft = recognized["drafts"][0]
    draft["review_status"] = "approved"
    confirm_response = await client.post(
        f"/api/papers/import/sessions/{recognized['session_id']}/confirm",
        json={
            "title": "导入历史试卷",
            "description": None,
            "root_knowledge_point_id": None,
            "drafts": [draft],
        },
    )
    assert confirm_response.status_code == 201
    payload = confirm_response.json()
    assert payload["title"] == "导入历史试卷"
    assert payload["source_type"] == "import"
    assert payload["question_count"] == 1

    session_response = await client.get(f"/api/papers/import/sessions/{recognized['session_id']}")
    assert session_response.status_code == 200
    assert session_response.json()["created_paper_id"] == payload["id"]
    assert session_response.json()["error_detail"] is None


@pytest.mark.asyncio
async def test_paper_import_recognize_file_extracts_docx_tables(client, db_session, monkeypatch):
    teacher = await _teacher(db_session)
    client.headers.update({"Authorization": f"Bearer {create_access_token(teacher.id, '')}"})
    prompts: list[str] = []

    async def fake_request(prompt: str) -> list[dict]:
        prompts.append(prompt)
        return [
            {
                "type": "choice",
                "content_text": "下面关于数据分析说法正确的是？",
                "options": {"A": "只做统计", "B": "服务决策", "C": "无需清洗", "D": "不能可视化"},
                "answer_text": "B",
                "analysis": "",
                "difficulty": 3,
                "raw_text": "下面关于数据分析说法正确的是？",
                "images": [],
            }
        ]

    monkeypatch.setattr("app.questions.service._request_doc_recognition_questions", fake_request)

    doc = Document()
    doc.add_paragraph("江苏卫生健康职业学院 2024-2025 学年试卷")
    table = doc.add_table(rows=5, cols=2)
    rows = [
        ("题号", "1"),
        ("题干", "下面关于数据分析说法正确的是？"),
        ("选项", "A. 只做统计\nB. 服务决策\nC. 无需清洗\nD. 不能可视化"),
        ("答案", "B"),
        ("解析", "数据分析用于支持业务决策。"),
    ]
    for index, (label, value) in enumerate(rows):
        table.rows[index].cells[0].text = label
        table.rows[index].cells[1].text = value
    buffer = BytesIO()
    doc.save(buffer)

    response = await client.post(
        "/api/papers/import/recognize-file",
        files={
            "file": (
                "数据分析试卷.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"prompt": "请严格按试卷题目识别，不要把表头当题目。"},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["session_id"]
    assert payload["summary"]["total"] == 1
    assert payload["drafts"][0]["content_text"] == "下面关于数据分析说法正确的是？"
    assert prompts and "[Q]" in prompts[0]


@pytest.mark.asyncio
async def test_paper_import_recognize_file_uses_pdf_chunked_recognizer(client, db_session, monkeypatch):
    teacher = await _teacher(db_session)
    client.headers.update({"Authorization": f"Bearer {create_access_token(teacher.id, '')}"})
    calls: list[tuple[bytes, str]] = []

    async def fake_pdf_recognizer(
        file_bytes: bytes,
        file_name: str,
        recognition_prompt: str | None = None,
    ) -> QuestionImportDocumentRecognizeResponse:
        calls.append((file_bytes, file_name))
        assert recognition_prompt is None
        drafts = [
            QuestionImportDraft(
                draft_id="choice-1",
                raw_text="1. SQL 中用于查询的关键字是？",
                title="SQL 中用于查询的关键字是？",
                type="choice",
                content_text="SQL 中用于查询的关键字是？",
                options={"A": "SELECT", "B": "UPDATE"},
                answer_text="A",
                segment_source="pdf_chunked_ai",
                type_confidence="high",
                boundary_confidence="high",
            ),
            QuestionImportDraft(
                draft_id="fill-1",
                raw_text="26. SQL 语言中，数据查询语句是 _____。",
                title="SQL 语言中，数据查询语句是 _____。",
                type="fill_in",
                content_text="SQL 语言中，数据查询语句是 _____。",
                answer_text="SELECT",
                segment_source="pdf_chunked_ai",
                type_confidence="high",
                boundary_confidence="high",
                answer_images=[
                    {
                        "image_id": "answer-er.png",
                        "url": "/api/uploads/files/answer-er.png",
                        "order": 1,
                    }
                ],
            ),
        ]
        return QuestionImportDocumentRecognizeResponse(
            mode="smart",
            summary=QuestionImportDocumentSummary(
                total=2,
                high_confidence=2,
                medium_confidence=0,
                low_confidence=0,
                issue_count=0,
                pending_review=2,
                approved=0,
                skipped=0,
            ),
            drafts=drafts,
        )

    monkeypatch.setattr("app.papers.service.recognize_pdf_with_ai", fake_pdf_recognizer)

    response = await client.post(
        "/api/papers/import/recognize-file",
        files={"file": ("sql-paper.pdf", b"%PDF fake content", "application/pdf")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert calls == [(b"%PDF fake content", "sql-paper.pdf")]
    assert payload["summary"]["total"] == 2
    assert [draft["type"] for draft in payload["drafts"]] == ["choice", "fill_in"]
    assert payload["drafts"][1]["answer_images"][0]["url"] == "/api/uploads/files/answer-er.png"

    session = await db_session.get(PaperImportSession, uuid.UUID(payload["session_id"]))
    assert session is not None
    assert session.source_format == "pdf"
    assert session.preview_payload["summary"]["total"] == 2
    assert session.preview_payload["drafts"][1]["answer_images"][0]["image_id"] == "answer-er.png"


def test_paper_import_extracts_docx_tables_in_document_order_and_splits_glued_questions():
    """PDF 转 DOCX 的试卷常把“答案：… 难度层次：易 12. 下一题”压在同一段。

    这类粘连题号必须在新行开始，且表格内容要留在原位置，否则 200+ 题的大试卷
    会被边界扫描漏掉大半（线上曾出现 250 题只识别出 11 题）。
    """
    from app.papers.service import extract_paper_import_file_content

    doc = Document()
    doc.add_paragraph("一、单选题")
    doc.add_paragraph("1. 第一题题干？")
    doc.add_paragraph("A. 甲\nB. 乙")
    doc.add_paragraph("答案：A 知识点：测试 难度层次：易 2. 第二题题干？")
    doc.add_paragraph("A. 丙\nB. 丁")
    doc.add_paragraph("答案：B 知识点：测试 难度层次：易")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "题干"
    table.rows[0].cells[1].text = "3. 第三题题干？"
    table.rows[1].cells[0].text = "答案"
    table.rows[1].cells[1].text = "A"
    doc.add_paragraph("4. 第四题题干？")
    buffer = BytesIO()
    doc.save(buffer)

    text, source_format = extract_paper_import_file_content("paper.docx", buffer.getvalue())

    assert source_format == "docx"
    # 粘连题号被拆到新行
    assert "\n2. 第二题题干？" in text
    # 表格内容按文档顺序出现在第二题与第四题之间
    assert text.index("2. 第二题题干？") < text.index("3. 第三题题干？") < text.index("4. 第四题题干？")


@pytest.mark.asyncio
async def test_paper_import_chunks_large_documents_instead_of_one_truncated_call(monkeypatch):
    """整卷一次性调用会把 200+ 题的输出截断；试卷导入必须分块识别。"""
    from app.questions.schemas import QuestionImportDocumentRecognizeRequest
    from app.questions.service import recognize_question_document

    prompts: list[str] = []
    seen = 0

    async def fake_request(prompt: str) -> list[dict]:
        nonlocal seen
        prompts.append(prompt)
        # 模板说明里也会出现 `[Q]` 字样，只统计「文本：」之后的文档正文
        body = prompt.rsplit("文本：", 1)[-1]
        question_count = body.count("[Q]")
        questions: list[dict] = []
        for _ in range(question_count):
            seen += 1
            questions.append(
                {
                    "type": "choice",
                    "content_text": f"第 {seen} 题题干？",
                    "options": {"A": "甲", "B": "乙", "C": "丙", "D": "丁"},
                    "answer_text": "A",
                    "analysis": "",
                    "difficulty": 3,
                    "raw_text": f"第 {seen} 题题干？",
                    "images": [],
                }
            )
        return questions

    monkeypatch.setattr("app.questions.service._request_doc_recognition_questions", fake_request)

    blocks = [
        f"{index}. 这是第 {index} 道题目，请选择正确答案？\nA. 甲选项内容\nB. 乙选项内容\nC. 丙选项内容\nD. 丁选项内容\n答案：A\n知识点：测试\n难度层次：易"
        for index in range(1, 121)
    ]
    raw_text = "一 、单选题(共120题)\n" + "\n".join(blocks)

    response = await recognize_question_document(
        QuestionImportDocumentRecognizeRequest(
            file_name="large-paper.docx",
            source_format="docx",
            raw_text=raw_text,
            analysis_mode="ai_full",
            import_context="paper",
        )
    )

    assert len(prompts) > 1, "大试卷必须分块调用 AI"
    assert "[Q]" in prompts[0]
    assert len(response.drafts) == 120


@pytest.mark.asyncio
async def test_paper_import_keeps_vision_path_when_page_images_are_present(monkeypatch):
    """带页面图片的试卷（PDF 图片回退）必须继续走视觉识别，而不是文本分块。"""
    from app.questions.schemas import QuestionImportDocumentRecognizeRequest
    from app.questions.service import recognize_question_document

    vision_calls: list[str] = []
    text_calls: list[str] = []

    async def fake_vision(**kwargs) -> dict:
        vision_calls.append(kwargs["prompt"])
        return {
            "questions": [
                {
                    "type": "choice",
                    "content_text": "图片中的题目？",
                    "options": {"A": "甲", "B": "乙"},
                    "answer_text": "A",
                    "analysis": "",
                    "difficulty": 3,
                    "raw_text": "图片中的题目？",
                    "images": [],
                }
            ]
        }

    async def fake_text(prompt: str) -> list[dict]:
        text_calls.append(prompt)
        return []

    monkeypatch.setattr("app.questions.service._request_vision_json", fake_vision)
    monkeypatch.setattr("app.questions.service._request_doc_recognition_questions", fake_text)

    response = await recognize_question_document(
        QuestionImportDocumentRecognizeRequest(
            file_name="paper.pdf",
            source_format="pdf",
            raw_text="[IMAGE:page-1]",
            analysis_mode="ai_full",
            import_context="paper",
            images=[
                {
                    "image_id": "page-1",
                    "url": "data:image/jpeg;base64,AAAA",
                    "order": 1,
                    "page": 1,
                }
            ],
        )
    )

    assert vision_calls, "带图片的试卷必须调用视觉识别"
    assert not text_calls, "带图片的试卷不应走文本分块识别"
    assert response.drafts[0].content_text == "图片中的题目？"


@pytest.mark.asyncio
async def test_paper_import_docx_uses_shared_question_bank_recognizer(client, db_session, monkeypatch):
    """试卷 DOCX 导入必须复用题库导入的识别器，两个页面结果保持一致。"""
    teacher = await _teacher(db_session)
    client.headers.update({"Authorization": f"Bearer {create_access_token(teacher.id, '')}"})
    calls: list[dict] = []

    async def fake_recognizer(file_bytes, file_name, *, import_context=None, recognition_prompt=None):
        calls.append(
            {
                "file_name": file_name,
                "import_context": import_context,
                "recognition_prompt": recognition_prompt,
            }
        )
        return QuestionImportDocumentRecognizeResponse(
            mode="smart",
            summary=QuestionImportDocumentSummary(
                total=1,
                high_confidence=1,
                medium_confidence=0,
                low_confidence=0,
                issue_count=0,
                pending_review=1,
                approved=0,
                skipped=0,
            ),
            drafts=[
                QuestionImportDraft(
                    draft_id=str(uuid.uuid4()),
                    raw_text="共享识别器题目？",
                    title="共享识别器题目？",
                    type="choice",
                    content_text="共享识别器题目？",
                    options={"A": "甲", "B": "乙"},
                    answer_text="A",
                    segment_source="ai_full",
                    boundary_confidence="high",
                    type_confidence="high",
                )
            ],
        )

    monkeypatch.setattr("app.questions.service.recognize_docx_with_ai", fake_recognizer)

    doc = Document()
    doc.add_paragraph("1. 共享识别器题目？")
    buffer = BytesIO()
    doc.save(buffer)

    response = await client.post(
        "/api/papers/import/recognize-file",
        files={
            "file": (
                "paper.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    assert calls == [
        {
            "file_name": "paper.docx",
            "import_context": "paper",
            "recognition_prompt": None,
        }
    ]
    assert response.json()["summary"]["total"] == 1


@pytest.mark.asyncio
async def test_paper_import_docx_falls_back_to_rule_drafts_when_ai_is_unavailable(
    client, db_session, monkeypatch
):
    """AI 不可用时保留原有回退：用规则解析生成草稿而不是直接报错。"""
    teacher = await _teacher(db_session)
    client.headers.update({"Authorization": f"Bearer {create_access_token(teacher.id, '')}"})

    async def broken_recognizer(*_args, **_kwargs):
        raise RuntimeError("AI 题目识别失败：error:RuntimeError:未配置任何 AI 识别服务的 API Key")

    monkeypatch.setattr("app.questions.service.recognize_docx_with_ai", broken_recognizer)

    doc = Document()
    doc.add_paragraph("一、单选题")
    doc.add_paragraph("1. 回退题目题干？")
    doc.add_paragraph("A. 甲")
    doc.add_paragraph("B. 乙")
    doc.add_paragraph("答案：A")
    buffer = BytesIO()
    doc.save(buffer)

    response = await client.post(
        "/api/papers/import/recognize-file",
        files={
            "file": (
                "paper.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["session_id"]
    assert payload["summary"]["total"] >= 1
