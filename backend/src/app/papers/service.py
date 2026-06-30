"""Core service functions for reusable paper assets."""

import random
from collections import Counter
from dataclasses import dataclass
from io import BytesIO
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pdfplumber
from docx import Document
from sqlalchemy import Select, and_, delete, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import joinedload, selectinload

from app.auth.models import User
from app.common.data_visibility import VisibilityScope
from app.common.resource_access import can_write_owned_resource, teacher_owned_resource_filter, teacher_visible_resource_filter
from app.learning.models import KnowledgePoint
from app.papers.models import Paper, PaperImportSession, PaperQuestion, PaperQuestionKnowledgeSuggestion
from app.papers.schemas import (
    PaperAIAppendRequest,
    PaperAIGenerateRequest,
    PaperCreate,
    PaperImportConfirmRequest,
    PaperImportRecognizeRequest,
    PaperQuestionItem,
    PaperUpdate,
)
from app.questions.ai_generate import AIGenerateRequest, AIModelProvider, generate_questions_stream
from app.questions.models import Question, QuestionBank
from app.questions.schemas import (
    ImportReviewStatus,
    QuestionCreate,
    QuestionImportAnalysisMode,
    QuestionImportDocumentRecognizeRequest,
    QuestionImportDocumentRecognizeResponse,
    QuestionImportDraft,
    QuestionImportImageInput,
)
from app.questions.similarity import (
    question_is_too_similar_to_any,
    question_summary_for_prompt,
)
from app.questions.service import (
    _request_deepseek_json,
    _load_root_descendant_knowledge_points,
    bulk_create_questions_fast,
    get_or_create_named_private_question_bank,
    recognize_pdf_with_ai,
    recognize_question_document,
)


@dataclass
class PaperGenerationProfile:
    total_count: int
    type_distribution: dict[str, int]
    difficulty: int
    knowledge_point_ids: list[uuid.UUID | str]
    prompt: str


@dataclass(frozen=True)
class GeneratedQuestionItemsResult:
    question_items: list[PaperQuestionItem]
    reused_source_question_ids: set[uuid.UUID]
    generated_question_count: int
    reused_source_question_count: int
    reused_bank_question_count: int


@dataclass(frozen=True)
class GenerationPlanSlot:
    source_item: Any
    question_type: str
    knowledge_point_ids: list[uuid.UUID]
    knowledge_point_names: list[str]
    knowledge_key: str
    knowledge_label: str
    type_distribution_text: str
    knowledge_distribution_text: str
    knowledge_quota_index: int
    knowledge_quota_count: int
    type_quota_index: int
    type_quota_count: int


@dataclass(frozen=True)
class GenerationPlan:
    slots: list[GenerationPlanSlot]
    type_distribution: dict[str, int]
    knowledge_distribution: dict[str, int]


def _format_docx_table(table: object, order: int) -> str:
    rows: list[list[str]] = []
    for row in table.rows:
        values = [cell.text.strip() for cell in row.cells]
        if any(values):
            rows.append(values)
    if not rows:
        return ""
    lines = [f"[TABLE:{order}]"]
    for values in rows:
        lines.append(" | ".join(values))
    return "\n".join(lines)


def extract_paper_import_file_content(file_name: str, file_bytes: bytes) -> tuple[str, str]:
    if not file_bytes:
        raise ValueError("文件内容为空")

    extension = Path(file_name).suffix.lower()
    if extension == ".markdown":
        extension = ".md"
    if extension not in {".pdf", ".docx", ".md"}:
        raise ValueError("暂不支持该文件格式，请上传 PDF、Word(docx) 或 Markdown 文件。")

    if extension == ".md":
        text = file_bytes.decode("utf-8-sig")
        if not text.strip():
            raise ValueError("文件内容为空")
        return text.strip(), "md"

    if extension == ".pdf":
        try:
            with pdfplumber.open(BytesIO(file_bytes)) as pdf:
                text = "\n\n".join(page.extract_text() or "" for page in pdf.pages)
        except Exception as exc:
            raise ValueError("PDF 文件解析失败，请检查文件后重试。") from exc
        if not text.strip():
            raise ValueError("未能从 PDF 中提取文字，请上传可复制文本的 PDF 或改用图片识别。")
        return text.strip(), "pdf"

    try:
        document = Document(BytesIO(file_bytes))
    except Exception as exc:
        raise ValueError("Word 文件解析失败，请检查文件后重试。") from exc

    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for index, table in enumerate(document.tables, start=1):
        table_text = _format_docx_table(table, index)
        if table_text:
            parts.append(table_text)
    text = "\n\n".join(parts).strip()
    if not text:
        raise ValueError("未能从 Word 文件中提取文字，请检查文件内容。")
    return text, "docx"


def build_paper_generation_profile(
    source_questions: list[dict],
    root_knowledge_point_id: uuid.UUID | str | None,
    difficulty_strategy: str,
    total_count: int | None = None,
) -> PaperGenerationProfile:
    if not source_questions:
        return PaperGenerationProfile(
            total_count=0,
            type_distribution={},
            difficulty=3,
            knowledge_point_ids=[root_knowledge_point_id] if root_knowledge_point_id else [],
            prompt="请参考源试卷结构生成一份新试卷。",
        )

    distribution = Counter(str(item.get("type") or "choice") for item in source_questions)
    target_total = total_count if total_count is not None else len(source_questions)
    difficulties = [int(item.get("difficulty") or 3) for item in source_questions]
    average = round(sum(difficulties) / len(difficulties)) if difficulties else 3
    if difficulty_strategy == "easier":
        average -= 1
    elif difficulty_strategy == "harder":
        average += 1
    difficulty = min(5, max(1, average))
    return PaperGenerationProfile(
        total_count=target_total,
        type_distribution=_scale_distribution(dict(distribution), target_total),
        difficulty=difficulty,
        knowledge_point_ids=[root_knowledge_point_id] if root_knowledge_point_id else [],
        prompt="请参考源试卷的题型结构、难度和考查范围，生成一份内容不同但能力要求接近的新试卷。",
    )


def _scale_distribution(distribution: dict[str, int], total_count: int) -> dict[str, int]:
    original_total = sum(distribution.values())
    if original_total <= 0 or total_count <= 0:
        return {}

    scaled: dict[str, int] = {}
    remainders: list[tuple[float, str]] = []
    for key, value in distribution.items():
        exact = total_count * value / original_total
        base = int(exact)
        if value > 0 and base == 0:
            base = 1
        scaled[key] = base
        remainders.append((exact - base, key))

    current = sum(scaled.values())
    if current < total_count:
        for _, key in sorted(remainders, reverse=True):
            if current >= total_count:
                break
            scaled[key] += 1
            current += 1
    elif current > total_count:
        for _, key in sorted(remainders):
            if current <= total_count:
                break
            if scaled[key] > 0:
                scaled[key] -= 1
                current -= 1

    return {key: value for key, value in scaled.items() if value > 0}


def _question_type_value(question: Question) -> str:
    return question.type.value if hasattr(question.type, "value") else str(question.type)


def _source_item_question_type(item: Any) -> str:
    return _question_type_value(item.question)


def _source_item_knowledge_point_ids(item: Any) -> list[uuid.UUID]:
    question = item.question
    return [knowledge_point.id for knowledge_point in question.knowledge_points or []]


def _source_item_knowledge_point_names(item: Any) -> list[str]:
    question = item.question
    return [
        str(knowledge_point.name).strip()
        for knowledge_point in question.knowledge_points or []
        if str(knowledge_point.name).strip()
    ]


def _source_item_primary_knowledge_key(
    item: Any,
    root_knowledge_point_id: uuid.UUID | None,
) -> str:
    knowledge_point_ids = _source_item_knowledge_point_ids(item)
    if knowledge_point_ids:
        return str(knowledge_point_ids[0])
    if root_knowledge_point_id is not None:
        return str(root_knowledge_point_id)
    return "__unmarked__"


def _source_item_generation_knowledge_ids(
    item: Any,
    root_knowledge_point_id: uuid.UUID | None,
) -> list[uuid.UUID]:
    knowledge_point_ids = _source_item_knowledge_point_ids(item)
    if knowledge_point_ids:
        return knowledge_point_ids
    return [root_knowledge_point_id] if root_knowledge_point_id is not None else []


def _source_item_generation_knowledge_label(
    item: Any,
    root_knowledge_point_id: uuid.UUID | None,
) -> tuple[str, list[str]]:
    names = _source_item_knowledge_point_names(item)
    if names:
        return " / ".join(names), names
    if root_knowledge_point_id is not None:
        return "课程主知识", ["课程主知识"]
    return "未标注知识点", []


def _source_item_score(item: Any) -> float | None:
    if item.score_override is not None:
        return float(item.score_override)
    if item.question is not None and item.question.score is not None:
        return float(item.question.score)
    return None


def _normalise_prompt_text(value: Any, *, max_length: int = 500) -> str:
    if value is None:
        return ""
    if isinstance(value, dict):
        preferred = value.get("text") or value.get("html") or value.get("markdown")
        if preferred is not None:
            value = preferred
    text = re.sub(r"\s+", " ", str(value)).strip()
    if len(text) > max_length:
        return f"{text[:max_length]}..."
    return text


def _format_source_question_context(slot: GenerationPlanSlot | Any) -> str:
    source_item = slot.source_item if isinstance(slot, GenerationPlanSlot) else slot
    question = source_item.question
    knowledge_point_names = "、".join(
        knowledge_point.name for knowledge_point in question.knowledge_points or []
    )
    parts = [
        f"源题题型：{_source_item_question_type(source_item)}",
        f"源题难度：{question.difficulty}",
        f"源题知识点：{knowledge_point_names or '未标注'}",
        f"源题标题：{_normalise_prompt_text(question.title, max_length=180)}",
        f"源题内容：{_normalise_prompt_text(question.content, max_length=520)}",
    ]
    if question.options:
        parts.append(f"源题选项：{_normalise_prompt_text(question.options, max_length=360)}")
    if question.answer:
        parts.append(f"源题答案：{_normalise_prompt_text(question.answer, max_length=260)}")
    if question.analysis:
        parts.append(f"源题解析：{_normalise_prompt_text(question.analysis, max_length=360)}")
    return _normalise_prompt_text("\n".join(part for part in parts if part), max_length=1400)


def _format_generation_plan_context(slot: GenerationPlanSlot | Any) -> str:
    if not isinstance(slot, GenerationPlanSlot):
        return ""
    return (
        "后端已完成本次模拟卷配额设计："
        f"整卷题型配额为 {slot.type_distribution_text}；"
        f"整卷知识点配额为 {slot.knowledge_distribution_text}；"
        f"当前题型 {slot.question_type} 为第 {slot.type_quota_index}/{slot.type_quota_count} 题；"
        f"当前知识点「{slot.knowledge_label}」为第 {slot.knowledge_quota_index}/{slot.knowledge_quota_count} 题。"
        "必须按照这个题型和知识点生成，不得自行切换到其它知识点或其它学科。"
    )


def _build_slot_ai_prompt(
    base_prompt: str,
    slot: GenerationPlanSlot | Any,
    selected_questions: list[Any] | None = None,
) -> str:
    source_context = _format_source_question_context(slot)
    plan_context = _format_generation_plan_context(slot)
    selected_questions = selected_questions or []
    selected_context = ""
    if selected_questions:
        summaries = [
            f"{index}. {question_summary_for_prompt(question)}"
            for index, question in enumerate(selected_questions[-8:], start=1)
        ]
        selected_context = "\n本卷已选题摘要（新题不得与这些题重复或近似）：\n" + "\n".join(summaries)
    prompt = f"""{base_prompt}

你正在为同一门课程、同一主知识点和同一子知识点补充一道新题。请先分析下方源题的学科领域、课程语境、知识点链路、考查能力和题型结构，再生成内容不同但考查目标等价的新题。
必须保持源题所属课程/主知识和子知识点，不得迁移到语文、英语、文学、历史、常识或其它无关学科；除非源题和知识点本身就是这些学科。
必须保持同一题型与相近难度。若源题是编程题，新题也必须是编程任务，并给出可判分参考答案。
严禁生成与本卷已选题、源题重复或基本近似的题目；不能只是替换变量名、数字、选项顺序或同义改写。
{plan_context}

源题参考：
{source_context}
{selected_context}
"""
    return _normalise_prompt_text(prompt, max_length=1900)


def _combined_generation_key(question_type: str, knowledge_key: str) -> str:
    return f"{question_type}\u241f{knowledge_key}"


def _split_combined_generation_key(key: str) -> tuple[str, str]:
    question_type, knowledge_key = key.split("\u241f", 1)
    return question_type, knowledge_key


def _build_generation_plan(
    source_items: list[Any],
    *,
    total_count: int,
    root_knowledge_point_id: uuid.UUID | None,
    prefer_root_knowledge_point: bool = False,
) -> GenerationPlan:
    grouped_items: dict[str, list[Any]] = {}
    key_order: list[str] = []
    knowledge_labels: dict[str, str] = {}

    for item in source_items:
        question_type = _source_item_question_type(item)
        if prefer_root_knowledge_point and root_knowledge_point_id is not None:
            knowledge_key = str(root_knowledge_point_id)
            label = "课程主知识"
        else:
            knowledge_key = _source_item_primary_knowledge_key(item, root_knowledge_point_id)
            label, _ = _source_item_generation_knowledge_label(item, root_knowledge_point_id)
        combined_key = _combined_generation_key(question_type, knowledge_key)
        if combined_key not in grouped_items:
            grouped_items[combined_key] = []
            key_order.append(combined_key)
        grouped_items[combined_key].append(item)
        knowledge_labels.setdefault(knowledge_key, label)

    combined_distribution = {
        key: len(items)
        for key, items in grouped_items.items()
    }
    combined_quota = _scale_distribution(combined_distribution, total_count)

    raw_slots: list[tuple[Any, str, str]] = []
    max_quota = max(combined_quota.values(), default=0)
    group_offsets: dict[str, int] = {}
    for round_index in range(max_quota):
        for key in key_order:
            if round_index >= combined_quota.get(key, 0):
                continue
            question_type, knowledge_key = _split_combined_generation_key(key)
            items = grouped_items[key]
            offset = group_offsets.get(key, 0)
            raw_slots.append((items[offset % len(items)], question_type, knowledge_key))
            group_offsets[key] = offset + 1

    type_distribution = Counter(question_type for _, question_type, _ in raw_slots)
    knowledge_distribution = Counter(knowledge_key for _, _, knowledge_key in raw_slots)
    type_distribution_text = "、".join(
        f"{question_type} {count}题"
        for question_type, count in type_distribution.items()
    ) or "无"
    knowledge_distribution_text = "、".join(
        f"{knowledge_labels.get(knowledge_key, knowledge_key)} {count}题"
        for knowledge_key, count in knowledge_distribution.items()
    ) or "无"
    type_seen: Counter[str] = Counter()
    knowledge_seen: Counter[str] = Counter()
    slots: list[GenerationPlanSlot] = []
    for item, question_type, knowledge_key in raw_slots:
        type_seen[question_type] += 1
        knowledge_seen[knowledge_key] += 1
        if prefer_root_knowledge_point and root_knowledge_point_id is not None:
            label = "课程主知识"
            names = ["课程主知识"]
            knowledge_point_ids = [root_knowledge_point_id]
        else:
            label, names = _source_item_generation_knowledge_label(item, root_knowledge_point_id)
            knowledge_point_ids = _source_item_generation_knowledge_ids(item, root_knowledge_point_id)
        slots.append(
            GenerationPlanSlot(
                source_item=item,
                question_type=question_type,
                knowledge_point_ids=knowledge_point_ids,
                knowledge_point_names=names,
                knowledge_key=knowledge_key,
                knowledge_label=knowledge_labels.get(knowledge_key, label),
                type_distribution_text=type_distribution_text,
                knowledge_distribution_text=knowledge_distribution_text,
                knowledge_quota_index=knowledge_seen[knowledge_key],
                knowledge_quota_count=knowledge_distribution[knowledge_key],
                type_quota_index=type_seen[question_type],
                type_quota_count=type_distribution[question_type],
            )
        )

    return GenerationPlan(
        slots=slots,
        type_distribution=dict(type_distribution),
        knowledge_distribution={
            knowledge_labels.get(key, key): count
            for key, count in knowledge_distribution.items()
        },
    )


def _question_scope_query_for_generation(*, user: User, is_admin: bool) -> Select:
    query = select(Question).where(Question.deleted_at.is_(None))
    if not is_admin:
        query = query.outerjoin(QuestionBank, Question.question_bank_id == QuestionBank.id).where(
            or_(
                teacher_owned_resource_filter(Question, user.id),
                and_(
                    Question.question_bank_id.is_not(None),
                    QuestionBank.visibility == VisibilityScope.PLATFORM,
                ),
            )
        )
    return query.options(selectinload(Question.knowledge_points))


async def _pick_existing_question_for_slot(
    db: AsyncSession,
    slot: GenerationPlanSlot,
    *,
    excluded_question_ids: set[uuid.UUID],
    selected_questions: list[Any],
    user: User,
    is_admin: bool,
    course_scope_kp_ids: set[uuid.UUID] | None = None,
) -> Question | None:
    source_question = slot.source_item.question
    type_value = slot.question_type
    knowledge_point_ids = slot.knowledge_point_ids
    base_query = _question_scope_query_for_generation(user=user, is_admin=is_admin).where(
        Question.type == type_value,
    )
    # 严格限定在本课程（主知识点及其子树）范围内取题，避免兜底时选到其它课程的无关题目，
    # 进而污染后续 AI 生成的上下文。
    if course_scope_kp_ids:
        base_query = base_query.where(
            Question.knowledge_points.any(KnowledgePoint.id.in_(course_scope_kp_ids))
        )
    if excluded_question_ids:
        base_query = base_query.where(Question.id.not_in(excluded_question_ids))

    query_attempts: list[Select] = []
    if source_question.question_bank_id is not None and knowledge_point_ids:
        query_attempts.append(
            base_query.where(
                Question.question_bank_id == source_question.question_bank_id,
                Question.knowledge_points.any(KnowledgePoint.id.in_(knowledge_point_ids)),
            )
        )
    if source_question.question_bank_id is not None:
        query_attempts.append(base_query.where(Question.question_bank_id == source_question.question_bank_id))
    if knowledge_point_ids:
        query_attempts.append(
            base_query.where(Question.knowledge_points.any(KnowledgePoint.id.in_(knowledge_point_ids)))
        )
    query_attempts.append(base_query)

    for query in query_attempts:
        rows = (await db.execute(query.limit(100))).scalars().unique().all()
        distinct_rows = [
            row for row in rows
            if not question_is_too_similar_to_any(row, [source_question, *selected_questions])
        ]
        if distinct_rows:
            return random.choice(list(distinct_rows))
    return None


async def generate_question_items_from_source_items(
    db: AsyncSession,
    source_items: list[Any],
    *,
    total_count: int,
    source_reuse_rate: int,
    root_knowledge_point_id: uuid.UUID | None,
    prefer_root_knowledge_point: bool,
    difficulty_strategy: str,
    model: str,
    user: User,
    is_admin: bool,
    exam_title: str | None = None,
) -> GeneratedQuestionItemsResult:
    source_items = [item for item in sorted(source_items, key=lambda value: value.order) if item.question is not None]
    if not source_items:
        raise ValueError("源内容没有可用于生成的题目")
    if total_count < len(source_items):
        raise ValueError("目标题目数不能少于源题目数")

    # 本课程（主知识点及其子树）的取题范围：复用题库题时严格限定在此范围内，
    # 避免选到其它课程的题目并污染 AI 生成上下文。
    course_scope_kp_ids: set[uuid.UUID] | None = None
    if root_knowledge_point_id is not None:
        root_uuid = (
            root_knowledge_point_id
            if isinstance(root_knowledge_point_id, uuid.UUID)
            else uuid.UUID(str(root_knowledge_point_id))
        )
        descendants = await _load_root_descendant_knowledge_points(db, root_uuid)
        course_scope_kp_ids = {root_uuid, *(kp.id for kp in descendants)}

    source_questions: list[dict] = [
        {
            "type": _source_item_question_type(item),
            "difficulty": item.question.difficulty,
            "knowledge_point_ids": _source_item_knowledge_point_ids(item),
        }
        for item in source_items
    ]
    profile = build_paper_generation_profile(
        source_questions,
        root_knowledge_point_id if prefer_root_knowledge_point else None,
        difficulty_strategy,
        total_count=total_count,
    )
    plan = _build_generation_plan(
        source_items,
        total_count=total_count,
        root_knowledge_point_id=root_knowledge_point_id,
        prefer_root_knowledge_point=prefer_root_knowledge_point,
    )
    slots = plan.slots
    if len(slots) != total_count:
        raise ValueError("无法按源题型和知识点比例生成题目计划")
    profile.type_distribution = plan.type_distribution

    allowed_source_reuse_count = min(
        len(source_items),
        int(total_count * source_reuse_rate / 100),
    )
    source_slots = list(range(min(len(source_items), len(slots))))
    reused_slot_indices = set(random.sample(source_slots, allowed_source_reuse_count)) if allowed_source_reuse_count > 0 else set()

    selected_question_ids: list[uuid.UUID | None] = [None] * len(slots)
    selected_scores: list[float | None] = [None] * len(slots)
    used_question_ids: set[uuid.UUID] = {item.question_id for item in source_items}
    selected_questions: list[Any] = []
    reused_source_question_ids: set[uuid.UUID] = set()
    reused_source_question_count = 0

    for index in reused_slot_indices:
        slot = slots[index]
        if question_is_too_similar_to_any(slot.source_item.question, selected_questions):
            continue
        selected_question_ids[index] = slot.source_item.question_id
        selected_scores[index] = _source_item_score(slot.source_item)
        selected_questions.append(slot.source_item.question)
        reused_source_question_ids.add(slot.source_item.question_id)
        reused_source_question_count += 1

    reused_bank_question_count = 0
    for index, slot in enumerate(slots):
        if selected_question_ids[index] is not None:
            continue
        picked = await _pick_existing_question_for_slot(
            db,
            slot,
            excluded_question_ids=used_question_ids,
            selected_questions=selected_questions,
            user=user,
            is_admin=is_admin,
            course_scope_kp_ids=course_scope_kp_ids,
        )
        if picked is None:
            continue
        selected_question_ids[index] = picked.id
        selected_scores[index] = _source_item_score(slot.source_item)
        used_question_ids.add(picked.id)
        selected_questions.append(picked)
        reused_bank_question_count += 1

    ai_slots = [index for index, question_id in enumerate(selected_question_ids) if question_id is None]
    generated_question_ids: list[uuid.UUID] = []
    generated_questions: list[QuestionCreate] = []
    for index in ai_slots:
        slot = slots[index]
        ai_profile = PaperGenerationProfile(
            total_count=1,
            type_distribution={slot.question_type: 1},
            difficulty=profile.difficulty,
            knowledge_point_ids=slot.knowledge_point_ids or profile.knowledge_point_ids,
            prompt=_build_slot_ai_prompt(profile.prompt, slot, selected_questions),
        )
        generated: QuestionCreate | None = None
        for retry_index in range(3):
            retry_prompt = ai_profile.prompt
            if retry_index > 0:
                retry_prompt = _normalise_prompt_text(
                    f"{ai_profile.prompt}\n前一次生成题与本卷已有题或源题过于相似，请换一个明显不同的任务、案例或考查角度。",
                    max_length=1900,
                )
            request = AIGenerateRequest(
                total_count=1,
                difficulty=ai_profile.difficulty,
                type_distribution=ai_profile.type_distribution,
                knowledge_point_ids=[kp for kp in ai_profile.knowledge_point_ids if isinstance(kp, uuid.UUID)],
                exam_title=exam_title or "",
                prompt=retry_prompt,
                model=AIModelProvider(model),
            )
            async for event in generate_questions_stream(db, request, user.id):
                event_type = str(event.get("type") or "")
                if event_type == "error":
                    raise ValueError(str(event.get("message") or "AI 生成失败"))
                if event_type != "question":
                    continue
                payload = event.get("data")
                if isinstance(payload, dict):
                    candidate = _question_create_from_ai_payload(payload, profile=ai_profile)
                    if question_is_too_similar_to_any(candidate, [slot.source_item.question, *selected_questions]):
                        break
                    generated = candidate
                    break
            if generated is not None:
                break
        if generated is None:
            raise ValueError("AI 生成题目与已有题过于相似，无法生成足够的不重复题目")
        source_bank_id = getattr(slot.source_item.question, "question_bank_id", None)
        if source_bank_id is not None:
            generated = generated.model_copy(update={"question_bank_id": source_bank_id})
        generated.score = _source_item_score(slot.source_item) or generated.score
        generated_questions.append(generated)
        selected_questions.append(generated)

    if generated_questions:
        result = await bulk_create_questions_fast(db, generated_questions, user.id)
        if len(result.created_question_ids) != len(generated_questions):
            raise ValueError("AI 生成题目入库数量不足")
        generated_question_ids = list(result.created_question_ids)

    for slot_index, question_id in zip(ai_slots, generated_question_ids, strict=True):
        selected_question_ids[slot_index] = question_id
        selected_scores[slot_index] = _source_item_score(slots[slot_index].source_item)

    question_items = [
        PaperQuestionItem(
            question_id=question_id,
            order=index,
            score_override=selected_scores[index],
        )
        for index, question_id in enumerate(selected_question_ids)
        if question_id is not None
    ]
    if len(question_items) != total_count:
        raise ValueError("生成题目数量不足")

    return GeneratedQuestionItemsResult(
        question_items=question_items,
        reused_source_question_ids=reused_source_question_ids,
        generated_question_count=len(generated_questions),
        reused_source_question_count=reused_source_question_count,
        reused_bank_question_count=reused_bank_question_count,
    )


def _question_create_from_ai_payload(
    payload: dict,
    *,
    profile: PaperGenerationProfile,
) -> QuestionCreate:
    raw_content = payload.get("content")
    if isinstance(raw_content, dict):
        content = raw_content
    elif isinstance(raw_content, str) and raw_content.strip():
        content = {"text": raw_content}
    else:
        content = {"text": str(payload.get("title") or "").strip()}

    raw_answer = payload.get("answer")
    if isinstance(raw_answer, dict):
        answer = raw_answer
    else:
        answer = {"text": str(raw_answer or "").strip()}

    raw_options = payload.get("options")
    options = raw_options if isinstance(raw_options, dict) else None

    difficulty_raw = payload.get("difficulty")
    difficulty = profile.difficulty
    if isinstance(difficulty_raw, (int, float)):
        difficulty = int(difficulty_raw)
    elif isinstance(difficulty_raw, str) and difficulty_raw.strip():
        try:
            difficulty = int(difficulty_raw.strip())
        except ValueError:
            difficulty = profile.difficulty
    difficulty = min(5, max(1, difficulty))

    knowledge_point_ids = [kp for kp in profile.knowledge_point_ids if isinstance(kp, uuid.UUID)]
    title = str(payload.get("title") or content.get("text") or "AI 生成题目").strip()[:500]
    if not title:
        title = "AI 生成题目"
    question_type = str(payload.get("type") or "choice")
    if question_type not in {"choice", "true_false", "fill_in", "short_answer", "essay", "code"}:
        question_type = "choice"

    return QuestionCreate(
        type=question_type,
        title=title,
        content=content,
        options=options,
        answer=answer,
        analysis=str(payload.get("analysis")).strip() if payload.get("analysis") is not None else None,
        difficulty=difficulty,
        score=10,
        knowledge_point_ids=knowledge_point_ids,
        tag_ids=[],
        question_bank_id=None,
    )


def paper_base_query() -> Select:
    question_load = selectinload(Paper.paper_questions).joinedload(PaperQuestion.question)
    return (
        select(Paper)
        .where(Paper.deleted_at.is_(None))
        .options(
            joinedload(Paper.creator),
            joinedload(Paper.root_knowledge_point),
            question_load,
            question_load.joinedload(Question.creator),
            question_load.joinedload(Question.question_bank),
            question_load.selectinload(Question.tags),
            question_load.selectinload(Question.knowledge_points),
        )
        .execution_options(populate_existing=True)
    )


def paper_scope_query(user: User, is_admin: bool) -> Select:
    query = paper_base_query()
    if not is_admin:
        query = query.where(teacher_owned_resource_filter(Paper, user.id))
    return query


async def list_papers_for_user(db: AsyncSession, *, user: User, is_admin: bool) -> tuple[list[Paper], int]:
    query = paper_scope_query(user, is_admin).order_by(Paper.created_at.desc())
    total = await db.scalar(select(func.count()).select_from(query.subquery()))
    rows = (await db.execute(query)).scalars().unique().all()
    papers = list(rows)
    for paper in papers:
        paper.paper_questions.sort(key=lambda item: item.order)
    return papers, int(total or 0)


async def get_paper_by_id(db: AsyncSession, paper_id: uuid.UUID, *, user: User, is_admin: bool) -> Paper | None:
    result = await db.execute(paper_scope_query(user, is_admin).where(Paper.id == paper_id))
    paper = result.scalars().unique().one_or_none()
    if paper:
        paper.paper_questions.sort(key=lambda item: item.order)
    return paper


async def list_paper_question_knowledge_suggestions(
    db: AsyncSession,
    paper_id: uuid.UUID,
) -> list[PaperQuestionKnowledgeSuggestion]:
    rows = await db.execute(
        select(PaperQuestionKnowledgeSuggestion)
        .where(
            PaperQuestionKnowledgeSuggestion.paper_id == paper_id,
            PaperQuestionKnowledgeSuggestion.deleted_at.is_(None),
        )
        .order_by(PaperQuestionKnowledgeSuggestion.created_at.asc())
    )
    return list(rows.scalars().all())


async def _list_all_paper_question_knowledge_suggestions(
    db: AsyncSession,
    paper_id: uuid.UUID,
) -> list[PaperQuestionKnowledgeSuggestion]:
    rows = await db.execute(
        select(PaperQuestionKnowledgeSuggestion)
        .where(PaperQuestionKnowledgeSuggestion.paper_id == paper_id)
        .order_by(PaperQuestionKnowledgeSuggestion.created_at.asc())
    )
    return list(rows.scalars().all())


async def _load_course_knowledge_point_scope(
    db: AsyncSession,
    root_knowledge_point_id: uuid.UUID | None,
) -> tuple[set[uuid.UUID], set[uuid.UUID]]:
    if root_knowledge_point_id is None:
        return set(), set()
    descendants = await _load_root_descendant_knowledge_points(db, root_knowledge_point_id)
    descendant_ids = {knowledge_point.id for knowledge_point in descendants}
    return {root_knowledge_point_id, *descendant_ids}, descendant_ids


def _question_needs_virtual_knowledge_suggestion(
    question: Question,
    *,
    course_knowledge_point_ids: set[uuid.UUID],
    concrete_knowledge_point_ids: set[uuid.UUID],
) -> bool:
    if not course_knowledge_point_ids:
        return False
    question_knowledge_point_ids = {knowledge_point.id for knowledge_point in question.knowledge_points or []}
    return not bool(question_knowledge_point_ids & concrete_knowledge_point_ids)


def _question_virtual_suggestion_prompt(question: Question, root_name: str | None) -> str:
    content = _normalise_prompt_text(question.content, max_length=1200)
    options = _normalise_prompt_text(question.options, max_length=600)
    answer = _normalise_prompt_text(question.answer, max_length=500)
    analysis = _normalise_prompt_text(question.analysis, max_length=500)
    return f"""
你是课程教研助手。当前课程没有维护足够的子知识点，系统不能真正关联不存在的知识点。
请根据题目内容，为它建议 1 个“虚拟知识点名称”，用于试卷知识点视角的临时覆盖展示。

课程/根知识点：{root_name or "当前课程"}
题型：{_question_type_value(question)}
题目标题：{question.title}
题目内容：{content}
选项：{options}
答案：{answer}
解析：{analysis}

要求：
1. suggested_name 用中文，2 到 12 个字，像课程目录里的知识点，不要包含“建议”“虚拟”“课程”等泛词。
2. reason 用一句话说明为什么这个题可能属于该知识点，最多 40 字。
3. 只返回合法 JSON，字段为 suggested_name: string, reason: string。
""".strip()


async def generate_paper_virtual_knowledge_suggestions(
    db: AsyncSession,
    paper: Paper,
    *,
    user_id: uuid.UUID,
) -> list[PaperQuestionKnowledgeSuggestion]:
    root_id = paper.root_knowledge_point_id
    course_knowledge_point_ids, concrete_knowledge_point_ids = await _load_course_knowledge_point_scope(db, root_id)
    existing = {
        suggestion.question_id: suggestion
        for suggestion in await _list_all_paper_question_knowledge_suggestions(db, paper.id)
    }
    candidates = [
        item
        for item in sorted(paper.paper_questions, key=lambda value: value.order)
        if item.question is not None
        and _question_needs_virtual_knowledge_suggestion(
            item.question,
            course_knowledge_point_ids=course_knowledge_point_ids,
            concrete_knowledge_point_ids=concrete_knowledge_point_ids,
        )
    ]
    candidate_question_ids = {item.question.id for item in candidates if item.question is not None}
    for question_id, suggestion in list(existing.items()):
        if question_id not in candidate_question_ids:
            suggestion.deleted_at = datetime.now(timezone.utc)
            existing.pop(question_id, None)

    for item in candidates:
        question = item.question
        if question.id in existing:
            existing[question.id].deleted_at = None
            continue

        fallback_name = _normalise_prompt_text(question.title or question.content, max_length=12) or "综合应用"
        fallback_name = re.sub(r"[^\u4e00-\u9fffA-Za-z0-9]", "", fallback_name)[:12] or "综合应用"
        suggested_name = fallback_name
        reason = "根据题干内容自动建议。"
        try:
            payload = await _request_deepseek_json(
                _question_virtual_suggestion_prompt(
                    question,
                    paper.root_knowledge_point.name if paper.root_knowledge_point else None,
                )
            )
            raw_name = str(payload.get("suggested_name") or "").strip()
            raw_reason = str(payload.get("reason") or "").strip()
            if raw_name:
                suggested_name = raw_name[:200]
            if raw_reason:
                reason = raw_reason[:200]
        except Exception:  # noqa: BLE001
            pass

        suggestion = PaperQuestionKnowledgeSuggestion(
            paper_id=paper.id,
            question_id=question.id,
            suggested_name=suggested_name,
            reason=reason,
            created_by=user_id,
        )
        db.add(suggestion)
        existing[question.id] = suggestion

    await db.flush()
    return await list_paper_question_knowledge_suggestions(db, paper.id)


def _normalized_paper_question_rows(paper_id: uuid.UUID, question_items: list[PaperQuestionItem]) -> list[PaperQuestion]:
    rows: list[PaperQuestion] = []
    seen_question_ids: set[uuid.UUID] = set()
    seen_orders: set[int] = set()
    for index, item in enumerate(question_items):
        if item.question_id in seen_question_ids:
            raise ValueError("duplicate question in paper")
        seen_question_ids.add(item.question_id)

        order = item.order if item.order is not None else index
        if order in seen_orders:
            raise ValueError("duplicate question order in paper")
        seen_orders.add(order)

        rows.append(
            PaperQuestion(
                paper_id=paper_id,
                question_id=item.question_id,
                order=order,
                score_override=item.score_override,
            )
        )
    return rows


async def _ensure_visible_source_paper(
    db: AsyncSession,
    source_paper_id: uuid.UUID | None,
    *,
    user: User,
    is_admin: bool,
) -> None:
    if source_paper_id is None:
        return
    query = select(Paper.id).where(Paper.id == source_paper_id, Paper.deleted_at.is_(None))
    if not is_admin:
        query = query.where(teacher_owned_resource_filter(Paper, user.id))
    if await db.scalar(query) is None:
        raise ValueError("source paper not found or not visible")


async def _ensure_visible_root_knowledge_point(
    db: AsyncSession,
    root_knowledge_point_id: uuid.UUID | None,
    *,
    user: User,
    is_admin: bool,
) -> None:
    if root_knowledge_point_id is None:
        return
    query = select(KnowledgePoint.id).where(
        KnowledgePoint.id == root_knowledge_point_id,
        KnowledgePoint.deleted_at.is_(None),
    )
    if not is_admin:
        query = query.where(teacher_visible_resource_filter(KnowledgePoint, user.id))
    if await db.scalar(query) is None:
        raise ValueError("knowledge point not found or not visible")


async def _ensure_visible_questions(
    db: AsyncSession,
    question_items: list[PaperQuestionItem],
    *,
    user: User,
    is_admin: bool,
) -> None:
    question_ids = {item.question_id for item in question_items}
    if not question_ids:
        return

    query = select(Question.id).where(Question.id.in_(question_ids), Question.deleted_at.is_(None))
    if not is_admin:
        query = query.outerjoin(QuestionBank, Question.question_bank_id == QuestionBank.id).where(
            or_(
                teacher_owned_resource_filter(Question, user.id),
                and_(
                    Question.question_bank_id.is_not(None),
                    QuestionBank.deleted_at.is_(None),
                    QuestionBank.visibility == VisibilityScope.PLATFORM,
                ),
            )
        )
    visible_ids = set(await db.scalars(query))
    if visible_ids != question_ids:
        raise ValueError("question not found or not visible")


async def _ensure_paper_references_visible(
    db: AsyncSession,
    *,
    user: User,
    is_admin: bool,
    source_paper_id: uuid.UUID | None = None,
    root_knowledge_point_id: uuid.UUID | None = None,
    question_items: list[PaperQuestionItem] | None = None,
) -> None:
    await _ensure_visible_source_paper(db, source_paper_id, user=user, is_admin=is_admin)
    await _ensure_visible_root_knowledge_point(db, root_knowledge_point_id, user=user, is_admin=is_admin)
    if question_items is not None:
        _normalized_paper_question_rows(uuid.uuid4(), question_items)
        await _ensure_visible_questions(db, question_items, user=user, is_admin=is_admin)


async def sync_paper_questions(db: AsyncSession, paper: Paper, question_items: list[PaperQuestionItem]) -> None:
    await db.execute(delete(PaperQuestion).where(PaperQuestion.paper_id == paper.id))
    await db.flush()
    rows = _normalized_paper_question_rows(paper.id, question_items)
    db.add_all(rows)


async def create_paper(db: AsyncSession, data: PaperCreate, *, user: User, is_admin: bool) -> Paper:
    await _ensure_paper_references_visible(
        db,
        user=user,
        is_admin=is_admin,
        source_paper_id=data.source_paper_id,
        root_knowledge_point_id=data.root_knowledge_point_id,
        question_items=data.question_items,
    )
    paper = Paper(
        title=data.title,
        description=data.description,
        source_type=data.source_type,
        source_paper_id=data.source_paper_id,
        root_knowledge_point_id=data.root_knowledge_point_id,
        is_reusable=data.is_reusable,
        created_by=user.id,
        owner_id=user.id,
    )
    db.add(paper)
    await db.flush()
    if data.question_items:
        await sync_paper_questions(db, paper, data.question_items)
    await db.flush()
    return (await get_paper_by_id(db, paper.id, user=user, is_admin=True)) or paper


async def update_paper(db: AsyncSession, paper: Paper, data: PaperUpdate, *, user: User, is_admin: bool) -> Paper:
    if not can_write_owned_resource(is_platform_admin=is_admin, current_user_id=user.id, owner_id=paper.owner_id):
        raise ValueError("paper not found or not writable")
    await _ensure_paper_references_visible(
        db,
        user=user,
        is_admin=is_admin,
        root_knowledge_point_id=data.root_knowledge_point_id if "root_knowledge_point_id" in data.model_fields_set else None,
        question_items=data.question_items,
    )
    values = data.model_dump(exclude_unset=True, exclude={"question_items"})
    for field, value in values.items():
        setattr(paper, field, value)
    if data.question_items is not None:
        await sync_paper_questions(db, paper, data.question_items)
    await db.flush()
    return (await get_paper_by_id(db, paper.id, user=user, is_admin=True)) or paper


async def archive_paper(db: AsyncSession, paper: Paper) -> Paper:
    paper.archived_at = datetime.now(timezone.utc)
    await db.flush()
    return paper


async def soft_delete_paper(db: AsyncSession, paper: Paper) -> None:
    paper.deleted_at = datetime.now(timezone.utc)
    await db.flush()


def _blocking_import_issues(draft: QuestionImportDraft) -> list[str]:
    return [issue for issue in draft.issues if not re.search(r"未识别到答案|缺少答案|缺答案", issue)]


def question_create_from_import_draft(
    draft: QuestionImportDraft,
    *,
    question_bank_id: uuid.UUID | None = None,
) -> QuestionCreate:
    answer_text = draft.answer_text or ""
    if draft.type.value == "choice":
        answer = {"correct": answer_text}
    elif draft.type.value == "true_false":
        answer = {"correct": answer_text.strip().lower() in {"正确", "对", "true", "t", "√"}}
    elif draft.type.value == "fill_in":
        answer = {"correct": [part.strip() for part in re.split(r"[;,；\n]", answer_text) if part.strip()]}
    elif draft.type.value == "code":
        answer = {"code": answer_text}
    else:
        answer = {"points": [part.strip() for part in answer_text.splitlines() if part.strip()]}
    if draft.answer_images:
        answer["images"] = [
            image.model_dump(mode="json") if hasattr(image, "model_dump") else image
            for image in draft.answer_images
        ]

    knowledge_point_ids = [
        kp.id for kp in (draft.suggested_knowledge_points or [])
    ]

    content_html = _build_import_content_html(draft.content_text, draft.images)

    return QuestionCreate(
        type=draft.type,
        title=(draft.title or draft.content_text[:120] or "未命名题目")[:500],
        content={"text": draft.content_text, "html": content_html},
        options=draft.options if draft.type.value == "choice" else None,
        answer=answer,
        analysis=draft.analysis,
        difficulty=draft.difficulty,
        score=10,
        knowledge_point_ids=knowledge_point_ids,
        tag_ids=[],
        question_bank_id=question_bank_id,
    )


def _escape_import_html(value: str) -> str:
    return (
        value.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
        .replace("'", "&#039;")
    )


def _build_import_content_html(
    content_text: str,
    images: list[QuestionImportImageInput] | None = None,
) -> str:
    lines = [line.strip() for line in content_text.splitlines() if line.strip()]
    parts: list[str] = []
    for line in lines:
        if line.startswith("<img "):
            parts.append(line)
        else:
            parts.append(f"<p>{_escape_import_html(line)}</p>")

    for image in images or []:
        alt = _escape_import_html((image.alt or "").strip() or "题目图片")
        src = _escape_import_html(image.url)
        image_id = _escape_import_html(image.image_id)
        parts.append(f'<img src="{src}" alt="{alt}" data-image-id="{image_id}" />')

    return "".join(parts)


async def create_import_session_from_recognition(
    db: AsyncSession,
    *,
    user: User,
    request: PaperImportRecognizeRequest,
) -> tuple[PaperImportSession, QuestionImportDocumentRecognizeResponse]:
    recognition = await recognize_question_document(
        QuestionImportDocumentRecognizeRequest(
            file_name=request.file_name,
            raw_text=request.raw_text,
            source_format=request.source_format,
            images=request.images,
            tables=request.tables,
            analysis_mode=QuestionImportAnalysisMode.AI_FULL,
            import_context="paper",
            recognition_prompt=request.recognition_prompt,
        )
    )
    session = await create_import_session_from_document_recognition(
        db,
        user=user,
        file_name=request.file_name,
        source_format=request.source_format,
        root_knowledge_point_id=request.root_knowledge_point_id,
        recognition=recognition,
    )
    return session, recognition


async def create_import_session_from_pdf_file(
    db: AsyncSession,
    *,
    user: User,
    file_name: str,
    file_bytes: bytes,
    root_knowledge_point_id: uuid.UUID | None = None,
    recognition_prompt: str | None = None,
) -> tuple[PaperImportSession, QuestionImportDocumentRecognizeResponse]:
    recognition = await recognize_pdf_with_ai(file_bytes, file_name, recognition_prompt=recognition_prompt)
    session = await create_import_session_from_document_recognition(
        db,
        user=user,
        file_name=file_name,
        source_format="pdf",
        root_knowledge_point_id=root_knowledge_point_id,
        recognition=recognition,
    )
    return session, recognition


async def create_import_session_from_document_recognition(
    db: AsyncSession,
    *,
    user: User,
    file_name: str,
    source_format: str,
    root_knowledge_point_id: uuid.UUID | None,
    recognition: QuestionImportDocumentRecognizeResponse,
) -> PaperImportSession:
    session = PaperImportSession(
        file_name=file_name,
        source_format=source_format,
        root_knowledge_point_id=root_knowledge_point_id,
        preview_payload=recognition.model_dump(mode="json"),
        error_detail=None,
        created_by=user.id,
    )
    db.add(session)
    await db.flush()
    return session


async def get_import_session(
    db: AsyncSession,
    session_id: uuid.UUID,
    *,
    user: User,
    is_admin: bool,
) -> PaperImportSession | None:
    stmt = select(PaperImportSession).where(
        PaperImportSession.id == session_id,
        PaperImportSession.deleted_at.is_(None),
    )
    if not is_admin:
        stmt = stmt.where(PaperImportSession.created_by == user.id)
    return (await db.execute(stmt)).scalar_one_or_none()


async def confirm_import_session(
    db: AsyncSession,
    session: PaperImportSession,
    body: PaperImportConfirmRequest,
    *,
    user: User,
    is_admin: bool,
) -> Paper:
    from app.questions.models import QuestionImportJob, QuestionImportJobStatus
    from app.questions.service import process_question_import_job

    root_id = body.root_knowledge_point_id or session.root_knowledge_point_id
    approved_drafts = [
        draft
        for draft in body.drafts
        if draft.review_status == ImportReviewStatus.APPROVED and not _blocking_import_issues(draft)
    ]
    if not approved_drafts:
        session.error_detail = "没有可入库的题目"
        await db.flush()
        raise ValueError("没有可入库的题目")

    question_bank = await get_or_create_named_private_question_bank(
        db,
        user_id=user.id,
        name=body.title,
        description=body.description or body.title,
    )
    questions = [
        question_create_from_import_draft(draft, question_bank_id=question_bank.id)
        for draft in approved_drafts
    ]
    try:
        result = await bulk_create_questions_fast(db, questions, user.id)
        if not result.created_question_ids and not result.existing_question_ids:
            session.error_detail = "没有可入库的题目"
            await db.flush()
            raise ValueError("没有可入库的题目")

        all_question_ids = list(result.created_question_ids) + list(result.existing_question_ids)
        all_scores = [q.score for q in result.created_questions] + [10] * len(result.existing_question_ids)

        paper = await create_paper(
            db,
            PaperCreate(
                title=body.title,
                description=body.description,
                source_type="import",
                root_knowledge_point_id=root_id,
                question_items=[
                    PaperQuestionItem(
                        question_id=question_id,
                        order=index,
                        score_override=all_scores[index],
                    )
                    for index, question_id in enumerate(all_question_ids)
                ],
            ),
            user=user,
            is_admin=is_admin,
        )
        session.created_paper_id = paper.id
        session.error_detail = None
        await db.flush()

        if root_id and all_question_ids and not body.skip_background_matching:
            # Only match KPs for NEW questions (existing ones already have KPs)
            new_question_ids = [str(qid) for qid in result.created_question_ids]
            job = QuestionImportJob(
                user_id=user.id,
                status=QuestionImportJobStatus.PENDING,
                total_count=len(result.created_question_ids),
                created_question_ids=new_question_ids,
            )
            db.add(job)
            await db.flush()
            await db.commit()
            import asyncio
            asyncio.create_task(
                process_question_import_job(
                    job_id=job.id,
                    user_id=user.id,
                    root_knowledge_point_id=root_id,
                    questions=[q.model_dump() for q in questions],
                )
            )

        return paper
    except Exception as exc:
        session.error_detail = str(exc)
        await db.flush()
        raise


async def generate_paper_from_source(
    db: AsyncSession,
    source: Paper,
    body: PaperAIGenerateRequest,
    *,
    user: User,
    is_admin: bool,
) -> Paper:
    source_items: list[PaperQuestion] = [
        item
        for item in sorted(source.paper_questions, key=lambda question_item: question_item.order)
        if item.question is not None
    ]
    if not source_items:
        raise ValueError("源试卷没有可用于生成的题目")

    generated = await generate_question_items_from_source_items(
        db,
        source_items,
        total_count=len(source_items),
        source_reuse_rate=body.source_reuse_rate,
        root_knowledge_point_id=source.root_knowledge_point_id,
        prefer_root_knowledge_point=body.prefer_root_knowledge_point,
        difficulty_strategy=body.difficulty_strategy,
        model=body.model,
        user=user,
        is_admin=is_admin,
        exam_title=source.title,
    )

    return await create_paper(
        db,
        PaperCreate(
            title=f"{source.title} - AI 生成",
            description=source.description,
            source_type="ai_generated",
            source_paper_id=source.id,
            root_knowledge_point_id=source.root_knowledge_point_id,
            question_items=generated.question_items,
        ),
        user=user,
        is_admin=is_admin,
    )


async def append_ai_questions_to_paper(
    db: AsyncSession,
    paper: Paper,
    body: PaperAIAppendRequest,
    *,
    user: User,
    is_admin: bool,
) -> Paper:
    if not can_write_owned_resource(is_platform_admin=is_admin, current_user_id=user.id, owner_id=paper.owner_id):
        raise ValueError("paper not found or not writable")

    source_items: list[PaperQuestion] = [
        item
        for item in sorted(paper.paper_questions, key=lambda question_item: question_item.order)
        if item.question is not None
    ]
    if not source_items:
        raise ValueError("当前试卷没有可用于 AI 生成的源题目")

    existing_question_ids = {item.question_id for item in source_items}
    generated = await generate_question_items_from_source_items(
        db,
        source_items,
        total_count=len(source_items) + body.question_count,
        source_reuse_rate=100,
        root_knowledge_point_id=paper.root_knowledge_point_id,
        prefer_root_knowledge_point=body.prefer_root_knowledge_point,
        difficulty_strategy=body.difficulty_strategy,
        model=body.model,
        user=user,
        is_admin=is_admin,
        exam_title=paper.title,
    )

    append_items = [
        item
        for item in generated.question_items
        if item.question_id not in existing_question_ids
    ][: body.question_count]
    if not append_items:
        raise ValueError("AI 未生成可追加的新题目，请调整数量或难度后重试")

    next_items = [
        PaperQuestionItem(
            question_id=item.question_id,
            order=index,
            score_override=item.score_override,
        )
        for index, item in enumerate(source_items)
    ]
    offset = len(next_items)
    next_items.extend(
        PaperQuestionItem(
            question_id=item.question_id,
            order=offset + index,
            score_override=item.score_override,
        )
        for index, item in enumerate(append_items)
    )
    await sync_paper_questions(db, paper, next_items)
    await db.flush()
    return (await get_paper_by_id(db, paper.id, user=user, is_admin=True)) or paper
