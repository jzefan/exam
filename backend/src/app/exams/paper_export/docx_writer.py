"""Render an :class:`ExamPaper` to a .docx byte string via python-docx."""

from __future__ import annotations

import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.exams.paper_export.model import ExamPaper, PaperSection, paper_header_lines

_CN_FONT = "宋体"
_GRID_CHUNK = 10  # choice answer grid: columns per row block
# Answer-key emphasis: blue text (no background fill).
_ANSWER_COLOR = RGBColor(0x1D, 0x4E, 0xD8)
# Two-character indent for in-stem continuation lines (sub-requirement lists).
_SUBITEM_INDENT = Pt(21)
# Blank writing space (number of empty lines) left after a question in the
# no-answer variant, by question type.
_BLANK_LINES = {"short_answer": 4, "essay": 7, "code": 7}


def _style_run(run, *, size: float, bold: bool = False, color: RGBColor | None = None) -> None:
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = _CN_FONT
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), _CN_FONT)


def _line(
    doc,
    text: str,
    *,
    size: float,
    bold: bool = False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    space_after: float = 4.0,
    color: RGBColor | None = None,
):
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.space_before = Pt(0)
    if text:
        _style_run(paragraph.add_run(text), size=size, bold=bold, color=color)
    return paragraph


def _set_cell(
    cell,
    text: str,
    *,
    size: float = 10.5,
    bold: bool = False,
    align=WD_ALIGN_PARAGRAPH.CENTER,
    color: RGBColor | None = None,
) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    if cell.text:
        paragraph.clear()
    _style_run(paragraph.add_run(text), size=size, bold=bold, color=color)


def _add_page_footer(doc) -> None:
    paragraph = doc.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def field(instr: str) -> None:
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr_el = OxmlElement("w:instrText")
        instr_el.set(qn("xml:space"), "preserve")
        instr_el.text = instr
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run = paragraph.add_run()
        run._element.append(begin)
        run2 = paragraph.add_run()
        run2._element.append(instr_el)
        run3 = paragraph.add_run()
        run3._element.append(end)

    _style_run(paragraph.add_run("第 "), size=9)
    field("PAGE")
    _style_run(paragraph.add_run(" 页 / 共 "), size=9)
    field("NUMPAGES")
    _style_run(paragraph.add_run(" 页"), size=9)


def _diagonal_cell(cell, top_right: str, bottom_left: str) -> None:
    """Header cell split by a top-left→bottom-right diagonal: ``top_right`` labels
    the columns to the right, ``bottom_left`` labels the row beneath."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    diagonal = OxmlElement("w:tl2br")
    diagonal.set(qn("w:val"), "single")
    diagonal.set(qn("w:sz"), "6")
    diagonal.set(qn("w:color"), "000000")
    borders.append(diagonal)

    top = cell.paragraphs[0]
    top.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    top.paragraph_format.space_after = Pt(0)
    top.paragraph_format.space_before = Pt(0)
    _style_run(top.add_run(top_right), size=9)
    bottom = cell.add_paragraph()
    bottom.alignment = WD_ALIGN_PARAGRAPH.LEFT
    bottom.paragraph_format.space_before = Pt(0)
    bottom.paragraph_format.space_after = Pt(0)
    _style_run(bottom.add_run(bottom_left), size=9)


def _score_table(doc, sections: tuple[PaperSection, ...]) -> None:
    """得分统计表: a diagonal 题号/得分 corner, one column per section, a 核查人签名
    column on the right, and an 阅卷教师 row at the bottom."""
    n = len(sections)
    cols = 1 + n + 1  # corner + sections + 核查人签名
    table = doc.add_table(rows=3, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    _diagonal_cell(table.rows[0].cells[0], "题号", "得分")
    for c, section in enumerate(sections, start=1):
        _set_cell(table.rows[0].cells[c], section.cn_index)
    _set_cell(table.rows[0].cells[cols - 1], "核查人签名", size=9)
    _set_cell(table.rows[2].cells[0], "阅卷教师", size=9)
    for row in table.rows:
        row.cells[0].width = Pt(54)
        row.cells[cols - 1].width = Pt(66)


def _score_box(doc) -> None:
    """Small 得分 box (得分 | blank) shown under each section heading."""
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    _set_cell(table.rows[0].cells[0], "得分", size=9)
    _set_cell(table.rows[0].cells[1], "", size=9)
    table.rows[0].cells[0].width = Pt(36)
    table.rows[0].cells[1].width = Pt(44)


def _choice_answer_grid(doc, section: PaperSection) -> None:
    questions = section.questions
    _line(doc, "选择题答案请填写下表中：", size=10.5, space_after=2)
    for start in range(0, len(questions), _GRID_CHUNK):
        chunk = questions[start : start + _GRID_CHUNK]
        table = doc.add_table(rows=2, cols=len(chunk))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        for idx, question in enumerate(chunk):
            _set_cell(table.rows[0].cells[idx], str(question.number), size=10)
            has_answer = bool(question.answer_text)
            _set_cell(
                table.rows[1].cells[idx],
                question.answer_text or "",
                size=10,
                color=_ANSWER_COLOR if has_answer else None,
            )


def _render_question(doc, question, *, with_answers: bool) -> None:
    # Per-question scores live in the section heading (每小题X分…), so the stem
    # carries only the number — no inline （X分）.
    stem_lines = question.stem.split("\n")
    first = stem_lines[0] if stem_lines else ""
    head = f"{question.number}. {first}".rstrip()
    para = _line(doc, head, size=10.5, space_after=2)
    para.paragraph_format.space_before = Pt(4)

    # Continuation lines (sub-requirement list / extra description) are indented so
    # their numbering reads as nested under the question's own number.
    for extra in stem_lines[1:]:
        sub = _line(doc, extra, size=10.5, space_after=1)
        sub.paragraph_format.left_indent = _SUBITEM_INDENT

    for option in question.options:
        _line(doc, f"    {option.label}. {option.text}", size=10.5, space_after=1)

    # Answer-key variant annotates every question (including choice) inline
    # with blue text; the grid stays as a summary.
    if question.answer_text is not None:
        _line(
            doc,
            f"正确答案：{question.answer_text}",
            size=10.5,
            space_after=2,
            color=_ANSWER_COLOR,
        )
    elif not with_answers:
        # Leave blank lines for the student to write 简答/论述/编程 answers.
        for _ in range(_BLANK_LINES.get(question.type, 0)):
            _line(doc, "", size=10.5, space_after=2)


def render_docx(paper: ExamPaper) -> bytes:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = _CN_FONT
    normal.font.size = Pt(10.5)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), _CN_FONT)

    # --- header block ---
    # Line 1: 学校名称 + 学期; Line 2: 《课程名》考试类型试卷; Line 3: （专业）.
    line1, line2, line3 = paper_header_lines(paper)
    _line(doc, line1, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _line(doc, line2, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _line(doc, line3, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)

    # Line 4: time limit + exam form.
    info = f"答题时限：{paper.duration_minutes} 分钟    考试形式：{paper.exam_form}"
    _line(doc, info, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

    # Line 5: class / id / name / score, centered.
    _line(
        doc,
        "班级__________ 学号__________ 姓名__________ 得分__________",
        size=12,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=6,
    )

    # --- score statistics table ---
    if paper.sections:
        _line(doc, "得分统计表：", size=10.5, space_after=2)
        _score_table(doc, paper.sections)
        _line(doc, "", size=6, space_after=2)

    # --- sections ---
    for section in paper.sections:
        _line(
            doc,
            f"{section.cn_index}、{section.title}（{section.score_note}）",
            size=12,
            bold=True,
            space_after=2,
        )
        _score_box(doc)
        if section.is_choice and section.questions:
            _choice_answer_grid(doc, section)
        for question in section.questions:
            _render_question(doc, question, with_answers=paper.with_answers)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
